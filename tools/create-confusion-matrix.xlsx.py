from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from pathlib import Path

out=Path(r'D:\โฟลเดอร์ใหม่ (6)\Confusion_Matrix_KitchenMart.xlsx')
wb=Workbook();ws=wb.active;ws.title='Confusion Matrix'
tp,tn,fp,fn=540,2375,25,60
ws.merge_cells('A1:D1');ws['A1']='Confusion Matrix — KitchenMart Behavior Detection';ws['A1'].font=Font(bold=True,size=16,color='FFFFFF');ws['A1'].fill=PatternFill('solid',fgColor='173F32');ws['A1'].alignment=Alignment(horizontal='center')
data=[['Prediction / Actual','Actual Normal','Actual Suspicious','รวม Predicted'],['Predicted Normal',tn,fn,tn+fn],['Predicted Suspicious',fp,tp,fp+tp],['รวม Actual',tn+fp,fn+tp,tn+fp+fn+tp]]
for r,row in enumerate(data,3):
    for c,val in enumerate(row,1):ws.cell(r,c,val)
for cell in ws[3]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245B49');cell.alignment=Alignment(horizontal='center')
for cell in ws['A']:
    if cell.row in (4,5,6):cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='E8F1EC')
colors={'B4':'D9EAD3','C4':'FCE5CD','B5':'F4CCCC','C5':'D9EAD3'}
for pos,color in colors.items():ws[pos].fill=PatternFill('solid',fgColor=color);ws[pos].font=Font(bold=True,size=13)
for row in ws.iter_rows(min_row=3,max_row=6,min_col=1,max_col=4):
    for cell in row:cell.alignment=Alignment(horizontal='center',vertical='center');cell.border=Border(*( [Side(style='thin',color='B7B7B7')]*4 ))
ws['A8']='ตัวชี้วัด';ws['B8']='สูตร';ws['C8']='ผลลัพธ์';ws['D8']='ร้อยละ'
metrics=[('Accuracy','(TP+TN)/(TP+TN+FP+FN)',(tp+tn)/(tp+tn+fp+fn)),('Precision','TP/(TP+FP)',tp/(tp+fp)),('Detection Rate / Recall','TP/(TP+FN)',tp/(tp+fn)),('F1-score','2×Precision×Recall/(Precision+Recall)',2*(tp/(tp+fp))*(tp/(tp+fn))/((tp/(tp+fp))+(tp/(tp+fn)))),('False Positive Rate','FP/(FP+TN)',fp/(fp+tn))]
for c in ws[8]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='E66B27');c.alignment=Alignment(horizontal='center')
for i,(name,formula,value) in enumerate(metrics,9):ws.cell(i,1,name);ws.cell(i,2,formula);ws.cell(i,3,value);ws.cell(i,3).number_format='0.0000';ws.cell(i,4,value);ws.cell(i,4).number_format='0.00%'
ws['A15']='คำอธิบาย';ws['A15'].font=Font(bold=True);notes=['TP = ระบบทำนาย Suspicious และเป็น Suspicious จริง','TN = ระบบทำนาย Normal และเป็น Normal จริง','FP = ระบบทำนาย Suspicious แต่เป็น Normal จริง','FN = ระบบทำนาย Normal แต่เป็น Suspicious จริง']
for i,n in enumerate(notes,16):ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=4);ws.cell(i,1,n)
ws.column_dimensions['A'].width=27;ws.column_dimensions['B'].width=38;ws.column_dimensions['C'].width=22;ws.column_dimensions['D'].width=20
for r in range(3,20):ws.row_dimensions[r].height=24
ws.freeze_panes='B4';ws.sheet_view.showGridLines=False;wb.save(out);print(out)
