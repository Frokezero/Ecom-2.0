from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path
import textwrap

source_path=Path(r'D:\โฟลเดอร์ใหม่ (6)\แบบเสนอเค้าโครงวิชาสัมมนา_KitchenMart.docx')
doc_path=Path(r'D:\โฟลเดอร์ใหม่ (6)\แบบเสนอเค้าโครงวิชาสัมมนา_KitchenMart_พร้อมโค้ด.docx')
root=Path(r'D:\ecommerce-php-app - สำเนา')
files=[
 ('สร้างข้อมูลผู้ใช้ 100 คน ระยะเวลา 30 วัน','tools/generate-behavior-dataset.php'),
 ('คำนวณ Mean และ Sample Standard Deviation','tools/analyze-behavior-baseline.php'),
 ('ทดลองตรวจจับและคำนวณ Confusion Matrix','tools/run-behavior-experiment.php'),
 ('ฟังก์ชันบันทึก Log และตรวจ Threshold ในระบบจริง','includes/behavior_analytics.php'),
]
d=Document(source_path);d.add_page_break()
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('ภาคผนวก ก\nโค้ดการทดลองระบบตรวจจับพฤติกรรมผิดปกติ');r.bold=True;r.font.name='TH Sarabun New';r._element.rPr.rFonts.set(qn('w:eastAsia'),'TH Sarabun New');r.font.size=Pt(18)
p=d.add_paragraph('ลำดับการทดลอง: สร้างข้อมูล → วิเคราะห์ Baseline → ตรวจ Threshold → สร้าง Confusion Matrix → คำนวณ Accuracy, Precision, Recall, F1-score, False Positive Rate, Detection Time และ System Performance');p.paragraph_format.first_line_indent=Cm(1.25)
p=d.add_paragraph();r=p.add_run('คำสั่งสำหรับทำการทดลองซ้ำ');r.bold=True;r.font.name='TH Sarabun New';r.font.size=Pt(16)
commands='php tools/migrate.php\nphp tools/seed-simulated-users.php\nphp tools/generate-behavior-dataset.php\nphp tools/analyze-behavior-baseline.php\nphp tools/run-behavior-experiment.php'
p=d.add_paragraph();r=p.add_run(commands);r.font.name='Consolas';r.font.size=Pt(9)
for title,rel in files:
    d.add_page_break();p=d.add_paragraph();r=p.add_run(title);r.bold=True;r.font.name='TH Sarabun New';r.font.size=Pt(16)
    p=d.add_paragraph();r=p.add_run(f'ไฟล์: {rel}');r.bold=True;r.font.name='Consolas';r.font.size=Pt(9)
    code=(root/rel).read_text(encoding='utf-8')
    for line in code.splitlines():
      for display_line in (textwrap.wrap(line,95,replace_whitespace=False,drop_whitespace=False) or [' ']):
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
        r=p.add_run(display_line);r.font.name='Consolas';r.font.size=Pt(7.5)
d.save(doc_path);print(doc_path)
